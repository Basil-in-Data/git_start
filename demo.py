import io
import os
import re
from datetime import datetime

import docx
import requests
import streamlit as st
from dotenv import load_dotenv
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from transformers import AutoModelForTokenClassification, AutoTokenizer, pipeline

# Official, stable Google GenAI client
from google import genai
from google.genai import types

GROQ_API_URL = "https://api.groq.com/openai/v1/chat/completions"
GROQ_MODEL = "llama-3.3-70b-versatile"

PII_MODEL_DIR = os.path.join(os.path.dirname(__file__), "pii_model")

# Each strategy is a reusable template with {{variable}} placeholders instead
# of one hardcoded topic slot. "defaults" both declares which variables the
# template needs (in display order) and pre-fills them, matching how
# enterprise prompt tools use named placeholders (e.g. {{audience}}) rather
# than free-typing the whole prompt from scratch.
PROMPT_TEMPLATES = {
    "Zero-Shot": {
        "template": "Explain the following topic clearly and concisely: {{topic}}",
        "defaults": {"topic": "Why data privacy matters in AI systems"},
    },
    "Few-Shot": {
        "template": (
            "Answer in exactly one sentence, following the style of these examples.\n\n"
            "Topic: Encryption\nAnswer: Encryption scrambles data so only authorized parties can read it.\n\n"
            "Topic: Firewalls\nAnswer: Firewalls filter network traffic to block unauthorized access.\n\n"
            "Topic: {{topic}}\nAnswer:"
        ),
        "defaults": {"topic": "Why data privacy matters in AI systems"},
    },
    "Role-Prompting": {
        "template": "You are a {{role}}. In a {{tone}} tone, explain: {{topic}}",
        "defaults": {
            "role": "senior data privacy officer briefing a company's executive team",
            "tone": "formal, authoritative",
            "topic": "Why data privacy matters in AI systems",
        },
    },
    "Chain-of-Thought (CoT)": {
        "template": (
            "Think through this step by step, showing your reasoning before giving "
            "a final answer.\n\nTopic: {{topic}}"
        ),
        "defaults": {"topic": "Why data privacy matters in AI systems"},
    },
}

# --- ENTERPRISE COLOR PALETTE (Ejada / ehub) ---
NAVY_DEEP = "#06154D"
NAVY_ACCENT = "#001C7F"
MINT = "#00E676"
MINT_ALT = "#00D27A"
PAGE_BG = "#F4F6FB"
CARD_BG = "#FFFFFF"
CARD_BORDER = "#E4E8F2"
TEXT_DARK = "#101B4D"
TEXT_MUTED = "#5B6478"
SIDEBAR_TEXT_MUTED = "#B7C0DE"


def extract_text_from_upload(uploaded_file) -> str:
    ext = uploaded_file.name.rsplit(".", 1)[-1].lower()
    raw = uploaded_file.read()
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(raw))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == "docx":
        document = docx.Document(io.BytesIO(raw))
        return "\n".join(p.text for p in document.paragraphs)
    return raw.decode("utf-8", errors="ignore")


def build_masked_document_pdf(clean_text, timestamp, findings, masking_purpose) -> bytes:
    """Builds the actual deliverable: a clean, shareable PDF of the redacted
    text meant to be handed off to another model or external party — not an
    internal audit log."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleNavy", parent=styles["Title"], textColor=HexColor(NAVY_DEEP), fontSize=18)
    heading_style = ParagraphStyle("HeadingNavy", parent=styles["Heading2"], textColor=HexColor(NAVY_ACCENT))
    body_style = ParagraphStyle("Body", parent=styles["BodyText"], textColor=HexColor("#1A1A1A"), fontSize=10, leading=14)

    type_counts = {}
    for f in findings:
        type_counts[f["type"]] = type_counts.get(f["type"], 0) + 1
    breakdown_lines = [f"{k}: {v} redacted" for k, v in type_counts.items()] or ["No PII detected"]

    story = [
        Paragraph("Redacted Document — Ready to Share", title_style),
        Spacer(1, 12),
        Paragraph(f"<b>Timestamp:</b> {timestamp}", body_style),
        Paragraph(f"<b>Masking Purpose:</b> {masking_purpose}", body_style),
        Paragraph(f"<b>Total Items Redacted:</b> {len(findings)}", body_style),
        Spacer(1, 8),
        Paragraph("Breakdown by PII Type:", heading_style),
    ]
    for line in breakdown_lines:
        story.append(Paragraph(f"— {line}", body_style))
    story.append(Spacer(1, 16))
    story.append(Paragraph("Masked Document:", heading_style))
    story.append(Spacer(1, 6))
    escaped = clean_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
    story.append(Paragraph(escaped, body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


# --- ENVIRONMENT & API SECURITY ---
# API keys are read only from the local .env file (GOOGLE_API_KEY, GROQ_API_KEY).
# They are never collected via the UI, so they can't leak into screenshots,
# logs, or session state.
load_dotenv()
GOOGLE_API_KEY = os.environ.get("GOOGLE_API_KEY")
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")

# --- PAGE LAYOUT & CONFIG ---
st.set_page_config(page_title="Ejada | Data Privacy Dashboard", page_icon="🔒", layout="wide")

# --- CUSTOM ENTERPRISE THEME (Navy sidebar / Mint accents / light workspace) ---
st.markdown(
    f"""
    <style>
    :root {{
        --navy-deep: {NAVY_DEEP};
        --navy-accent: {NAVY_ACCENT};
        --mint: {MINT};
        --mint-alt: {MINT_ALT};
        --page-bg: {PAGE_BG};
        --card-bg: {CARD_BG};
        --card-border: {CARD_BORDER};
        --text-dark: {TEXT_DARK};
        --text-muted: {TEXT_MUTED};
        --sidebar-text-muted: {SIDEBAR_TEXT_MUTED};
    }}

    .stApp {{ background: var(--page-bg); }}
    [data-testid="stHeader"] {{ background: rgba(0,0,0,0); }}
    [data-testid="stAppViewBlockContainer"] {{ padding-top: 0; }}

    /* --- Sidebar (kept navy, matches the reference mock's nav rail) --- */
    [data-testid="stSidebar"] {{
        background: var(--navy-deep);
    }}
    [data-testid="stSidebar"] * {{ color: var(--sidebar-text-muted); }}
    .sidebar-logo {{
        display: flex;
        align-items: center;
        gap: 0.6rem;
        font-size: 1.15rem;
        font-weight: 700;
        color: #FFFFFF;
        padding: 0.5rem 0 1.5rem 0;
        border-bottom: 1px solid rgba(255,255,255,0.12);
        margin-bottom: 1rem;
    }}
    .sidebar-logo-badge {{
        background: var(--mint);
        color: var(--navy-deep);
        width: 32px;
        height: 32px;
        border-radius: 8px;
        display: flex;
        align-items: center;
        justify-content: center;
        font-weight: 800;
    }}
    /* Real nav buttons (st.button in the sidebar) styled as nav items:
       full-width, left-aligned, active page shown via type="primary". */
    [data-testid="stSidebar"] div.stButton {{ margin-bottom: 0.25rem; }}
    [data-testid="stSidebar"] div.stButton > button {{
        width: 100%;
        justify-content: flex-start;
        text-align: left;
        background: transparent;
        border: none;
        border-radius: 8px;
        color: var(--sidebar-text-muted);
        font-weight: 500;
        padding: 0.6rem 0.75rem;
    }}
    [data-testid="stSidebar"] div.stButton > button:hover {{
        background: rgba(255,255,255,0.08);
        color: #FFFFFF;
    }}
    [data-testid="stSidebar"] div.stButton > button[kind="primary"] {{
        background: rgba(0,230,118,0.12) !important;
        border-left: 3px solid var(--mint) !important;
        color: #FFFFFF !important;
        font-weight: 700;
    }}
    .sidebar-footer {{
        margin-top: 2rem;
        padding-top: 1rem;
        border-top: 1px solid rgba(255,255,255,0.12);
        font-size: 0.85rem;
        color: var(--sidebar-text-muted);
    }}
    .sidebar-divider {{
        border-top: 1px solid rgba(255,255,255,0.12);
        margin: 0.5rem 0 0.75rem 0;
    }}

    /* --- Top banner (mint) --- */
    .top-banner {{
        display: flex;
        align-items: center;
        justify-content: space-between;
        background: linear-gradient(90deg, var(--mint), var(--mint-alt));
        color: var(--navy-deep);
        padding: 0.85rem 1.5rem;
        border-radius: 12px;
        margin-bottom: 1.25rem;
        font-weight: 600;
    }}
    .top-banner-icons {{ font-size: 1.1rem; letter-spacing: 0.5rem; }}

    .page-title {{
        font-size: 2rem;
        font-weight: 800;
        color: var(--text-dark);
        margin-bottom: 1.25rem;
    }}

    h1, h2, h3, h4 {{ color: var(--text-dark) !important; }}
    p, li, span, label {{ color: var(--text-muted); }}
    [data-testid="stSidebar"] h1, [data-testid="stSidebar"] h2,
    [data-testid="stSidebar"] h3, [data-testid="stSidebar"] h4 {{ color: #FFFFFF !important; }}

    /* Buttons carry their own text color; any inner element must inherit
       it, otherwise the broad paragraph text-color rule above wins (a
       rule directly on an element always beats an inherited value,
       regardless of the ancestor's specificity). */
    div.stButton > button *, div.stDownloadButton > button *, button[kind="secondary"] *,
    button[kind="primary"] * {{ color: inherit !important; }}

    /* --- Buttons (covers both st.button/st.download_button wrappers and
       Streamlit's own internal buttons, e.g. the file uploader's "Browse"
       button, which render as bare button[kind=...] with no .stButton
       wrapper around them) --- */
    div.stButton > button, div.stDownloadButton > button,
    button[kind="secondary"], button[kind="primary"] {{
        background: var(--card-bg);
        color: var(--navy-deep);
        border: 2px solid var(--navy-deep);
        border-radius: 999px;
        font-weight: 600;
        padding: 0.6rem 1.8rem;
        transition: all 0.15s ease;
    }}
    div.stButton > button:hover, div.stDownloadButton > button:hover,
    button[kind="secondary"]:hover {{
        background: var(--navy-deep);
        color: #FFFFFF;
    }}
    div.stButton > button[kind="primary"], button[kind="primary"] {{
        background: linear-gradient(90deg, var(--mint), var(--mint-alt));
        color: var(--navy-deep);
        border: none;
        font-weight: 700;
    }}
    div.stButton > button[kind="primary"]:hover, button[kind="primary"]:hover {{
        filter: brightness(1.05);
        color: var(--navy-deep);
        box-shadow: 0 4px 14px rgba(0,230,118,0.35);
    }}

    /* --- Cards (st.container(border=True)) --- */
    [data-testid="stVerticalBlockBorderWrapper"] {{
        background: var(--card-bg);
        border: 1px solid var(--card-border) !important;
        border-radius: 14px !important;
        padding: 0.5rem 0.25rem;
        box-shadow: 0 2px 10px rgba(16,27,77,0.06);
    }}
    .card-heading {{
        font-weight: 700;
        color: var(--text-dark);
        font-size: 1.05rem;
        margin-bottom: 0.75rem;
    }}

    /* --- Inputs --- */
    [data-testid="stTextArea"] textarea,
    [data-testid="stTextInput"] input,
    [data-testid="stFileUploaderDropzone"] {{
        background: #FBFCFE !important;
        color: var(--text-dark) !important;
        border: 1px solid var(--card-border) !important;
        border-radius: 8px !important;
    }}
    [data-testid="stTextArea"] textarea:focus,
    [data-testid="stTextInput"] input:focus {{
        border-color: var(--mint) !important;
        box-shadow: 0 0 0 1px var(--mint) !important;
    }}
    /* Disabled textareas (e.g. the read-only "Masked Output" box) get faded
       by the browser using -webkit-text-fill-color, which silently wins
       over a plain `color` rule even with !important — set it explicitly
       so the output text renders fully black, not washed out. */
    [data-testid="stTextArea"] textarea:disabled {{
        color: #000000 !important;
        -webkit-text-fill-color: #000000 !important;
        opacity: 1 !important;
    }}

    [data-testid="stSelectbox"] div[data-baseweb="select"] > div {{
        background: #FBFCFE !important;
        border: 1px solid var(--card-border) !important;
        color: var(--text-dark) !important;
        border-radius: 8px !important;
    }}

    [data-testid="stAlert"] {{
        border-radius: 10px !important;
        border-left: 4px solid var(--mint) !important;
        background: #FFFFFF !important;
    }}
    [data-testid="stAlert"] p {{ color: var(--text-dark) !important; }}

    [data-testid="stCodeBlock"] pre {{
        background: var(--navy-deep) !important;
        border: 1px solid var(--card-border) !important;
        border-radius: 8px !important;
        color: #E4F8ED !important;
    }}

    [data-testid="stCaptionContainer"] {{ color: var(--text-muted) !important; }}
    </style>
    """,
    unsafe_allow_html=True,
)

TRANSPARENCY_PAGE = ("How Your Data Is Masked", "🛡️")

NAV_PAGES = [
    ("Secure your files", "🔒"),
    ("Prompt Engineering", "💡"),
    ("Chat", "💬"),
    ("Request a Model", "🔌"),
]

if "active_page" not in st.session_state:
    st.session_state.active_page = "Secure your files"


def _sidebar_nav_button(page_name: str, icon: str):
    is_active = st.session_state.active_page == page_name
    if st.button(
        f"{icon}  {page_name}",
        key=f"nav_{page_name}",
        type="primary" if is_active else "secondary",
    ):
        st.session_state.active_page = page_name
        st.rerun()


with st.sidebar:
    st.markdown(
        """
        <div class="sidebar-logo">
            <div class="sidebar-logo-badge">e</div>
            <div>Ejada Workspace</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    # Kept separate from the main workflow nav below — this is a trust/
    # transparency page, not a fifth workflow step.
    _sidebar_nav_button(*TRANSPARENCY_PAGE)
    st.markdown('<div class="sidebar-divider"></div>', unsafe_allow_html=True)

    for page_name, icon in NAV_PAGES:
        _sidebar_nav_button(page_name, icon)

    st.markdown(
        """
        <div class="sidebar-footer">
            <b style="color:#FFFFFF;">ejada</b>
        </div>
        """,
        unsafe_allow_html=True,
    )

st.markdown(
    """
    <div class="top-banner">
        <span>Secure Enterprise AI Masking & Prompting Environment</span>
        <span class="top-banner-icons">💬 🔔 👤</span>
    </div>
    """,
    unsafe_allow_html=True,
)
st.markdown('<div class="page-title">The AI Gateway</div>', unsafe_allow_html=True)

if not GOOGLE_API_KEY:
    st.error(
        "No `GOOGLE_API_KEY` found. Create a `.env` file in the project root with:\n\n"
        "```\nGOOGLE_API_KEY=your_key_here\n```"
    )

client = genai.Client() if GOOGLE_API_KEY else None


def _call_gemini(prompt: str, system_prompt: str):
    config = types.GenerateContentConfig(system_instruction=system_prompt) if system_prompt.strip() else None
    response = client.models.generate_content(model="gemini-2.5-flash", contents=prompt, config=config)
    usage = None
    meta = getattr(response, "usage_metadata", None)
    if meta is not None:
        usage = {
            "prompt_tokens": meta.prompt_token_count,
            "completion_tokens": meta.candidates_token_count,
            "total_tokens": meta.total_token_count,
        }
    return response.text, "Gemini (gemini-2.5-flash)", usage


def _call_groq(prompt: str, system_prompt: str):
    messages = []
    if system_prompt.strip():
        messages.append({"role": "system", "content": system_prompt})
    messages.append({"role": "user", "content": prompt})
    resp = requests.post(
        GROQ_API_URL,
        headers={"Authorization": f"Bearer {GROQ_API_KEY}"},
        json={"model": GROQ_MODEL, "messages": messages},
        timeout=30,
    )
    resp.raise_for_status()
    data = resp.json()
    usage_raw = data.get("usage") or {}
    usage = (
        {
            "prompt_tokens": usage_raw.get("prompt_tokens"),
            "completion_tokens": usage_raw.get("completion_tokens"),
            "total_tokens": usage_raw.get("total_tokens"),
        }
        if usage_raw
        else None
    )
    return data["choices"][0]["message"]["content"], f"Groq ({GROQ_MODEL})", usage


def generate_with_fallback(prompt: str, system_prompt: str = "", preferred: str = "gemini"):
    """Try the user-selected backend first; on any error (or if it's not
    configured), fall back to the other backend if available.
    Returns (response_text, model_used, usage_dict_or_None)."""
    backends = [
        ("gemini", _call_gemini, client is not None),
        ("groq", _call_groq, bool(GROQ_API_KEY)),
    ]
    if preferred == "groq":
        backends.reverse()

    first_error = None
    for name, call_fn, available in backends:
        if not available:
            continue
        try:
            text, label, usage = call_fn(prompt, system_prompt)
            if first_error is not None:
                label = f"{label} — fallback (preferred backend failed: {first_error})"
            return text, label, usage
        except Exception as e:
            first_error = e

    if first_error is not None:
        raise first_error
    raise RuntimeError("No API key configured for Gemini or Groq.")


# ==========================================
# PII CLASSIFICATION & MASKING MODEL
# ==========================================
class PIIMaskingModel:
    """
    Wraps the fine-tuned BertForTokenClassification model in `pii_model/`
    (36 BIO labels: EMAIL, GIVENNAME, TELEPHONENUM, SOCIALNUM, etc.).
    detect(text) -> list of findings, mask(text) -> (clean_text, findings).
    """

    # Structurally-unambiguous PII formats where the token-classifier's raw
    # per-token BIO tags tend to leave gaps mid-entity (e.g. only the local
    # part of an email, or the area code of a phone number). Used as a
    # safety net so a real email/phone is never left partially exposed.
    SAFETY_PATTERNS = {
        "EMAIL": re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"),
        "CREDITCARDNUMBER": re.compile(r"\b(?:\d[ -]?){13,16}\d\b"),
        "TELEPHONENUM": re.compile(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b"),
        "SOCIALNUM": re.compile(r"\b\d{3}-\d{2}-\d{4}\b"),
    }

    def __init__(self, model_dir: str):
        tokenizer = AutoTokenizer.from_pretrained(model_dir)
        model = AutoModelForTokenClassification.from_pretrained(model_dir)
        self.nlp = pipeline(
            "token-classification",
            model=model,
            tokenizer=tokenizer,
            aggregation_strategy="simple",
        )

    @staticmethod
    def _overlaps(a_start, a_end, b_start, b_end):
        return not (a_end <= b_start or a_start >= b_end)

    @staticmethod
    def _snap_to_word_boundaries(text: str, findings: list):
        # BERT classifies subwords independently, so it sometimes tags only
        # the first WordPiece of a word (e.g. "B" of "Basil") and mislabels
        # the rest as non-PII — those continuation pieces never even reach
        # the merge step above since they're not part of any entity. Expand
        # every finding out to the surrounding whitespace so a partially
        # tagged word is never masked as just its first letter.
        for f in findings:
            start = f["start"]
            while start > 0 and not text[start - 1].isspace():
                start -= 1
            end = f["end"]
            while end < len(text) and not text[end].isspace():
                end += 1
            f["start"], f["end"] = start, end
        return findings

    @staticmethod
    def _merge_overlaps(findings: list):
        # Snapping to word boundaries can make two previously-separate
        # findings overlap (e.g. different subwords of the same word were
        # tagged with different types) — collapse those into one finding.
        findings = sorted(findings, key=lambda f: f["start"])
        merged = []
        for f in findings:
            if merged and f["start"] < merged[-1]["end"]:
                prev = merged[-1]
                prev["end"] = max(prev["end"], f["end"])
                if f["score"] > prev["score"]:
                    prev["type"] = f["type"]
                    prev["score"] = f["score"]
            else:
                merged.append(f)
        return merged

    def _apply_safety_net(self, text: str, findings: list):
        for pii_type, pattern in self.SAFETY_PATTERNS.items():
            for m in pattern.finditer(text):
                span_start, span_end = m.span()
                overlapping = [f for f in findings if self._overlaps(f["start"], f["end"], span_start, span_end)]
                if overlapping:
                    start = min([span_start] + [f["start"] for f in overlapping])
                    end = max([span_end] + [f["end"] for f in overlapping])
                    best_type = max(overlapping, key=lambda f: f["score"])["type"]
                    findings = [f for f in findings if f not in overlapping]
                    findings.append({"type": best_type, "start": start, "end": end, "score": 0.99})
                else:
                    findings.append({"type": pii_type, "start": span_start, "end": span_end, "score": 0.99})
        return sorted(findings, key=lambda f: f["start"])

    def detect(self, text: str, gap: int = 2):
        # The BIO tagger sometimes flips label within one PII instance on
        # punctuation/no-space runs (e.g. a phone number); bridge small gaps
        # between adjacent entity spans into a single finding.
        raw = sorted(self.nlp(text), key=lambda e: e["start"])
        merged = []
        for e in raw:
            if merged and e["start"] - merged[-1]["end"] <= gap:
                prev = merged[-1]
                prev["end"] = max(prev["end"], e["end"])
                if e["score"] > prev["score"]:
                    prev["type"] = e["entity_group"]
                    prev["score"] = e["score"]
            else:
                merged.append({
                    "type": e["entity_group"],
                    "start": e["start"],
                    "end": e["end"],
                    "score": e["score"],
                    "value": text[e["start"]:e["end"]],
                })
        merged = self._snap_to_word_boundaries(text, merged)
        merged = self._merge_overlaps(merged)
        merged = self._apply_safety_net(text, merged)
        for f in merged:
            f["value"] = text[f["start"]:f["end"]]
        return merged

    def mask(self, text: str, strict: bool = False):
        """strict=True (vendor sharing) uses a generic [REDACTED] tag that hides
        the PII category entirely. strict=False (LLM ingestion) keeps the typed
        tag (e.g. [REDACTED_EMAIL]) so a downstream model retains useful context."""
        findings = self.detect(text)
        clean_text = text
        for f in sorted(findings, key=lambda f: -f["start"]):
            tag = "[REDACTED]" if strict else f"[REDACTED_{f['type']}]"
            clean_text = clean_text[: f["start"]] + tag + clean_text[f["end"] :]
        return clean_text, findings


@st.cache_resource(show_spinner="Loading PII classification model...")
def load_pii_model():
    return PIIMaskingModel(PII_MODEL_DIR)


pii_model = load_pii_model()

# --- PAGE ROUTING (driven by the sidebar nav buttons above) ---
if st.session_state.active_page == "Secure your files":
    st.write(
        "Upload or paste a sensitive document, log the reason it needs to be shared, "
        "then run it through the PII masking pipeline before handing it off."
    )

    card_col1, card_col2 = st.columns(2)

    with card_col1:
        with st.container(border=True):
            st.markdown('<div class="card-heading">1. Document Redaction Input</div>', unsafe_allow_html=True)
            uploaded_file = st.file_uploader(
                "Upload a document (.txt, .csv, .log, .md, .json, .pdf, .docx)",
                type=["txt", "csv", "log", "md", "json", "pdf", "docx"],
            )
            pasted_text = st.text_area(
                "...or paste the document text directly",
                height=160,
                placeholder="e.g. Please contact Basil at basil@example.com or (555) 123-4567 regarding the ticket.",
            )

    with card_col2:
        with st.container(border=True):
            st.markdown('<div class="card-heading">2. Masking Configuration</div>', unsafe_allow_html=True)
            masking_purpose = st.selectbox(
                "Masking Purpose",
                [
                    "AI Tool Ingestion (Optimized for Large Language Models)",
                    "External Vendor Sharing (Strict Data Privacy Redaction)",
                ],
            )
            st.caption(
                "🔒 Masking always executes locally via the fine-tuned PII model below — "
                "no document text is ever sent to an external API."
            )
            run_pipeline = st.button("🚀 Run PII Masking Pipeline", type="primary")

    if run_pipeline:
        source_text = ""
        if uploaded_file is not None:
            source_text = extract_text_from_upload(uploaded_file)
        elif pasted_text.strip():
            source_text = pasted_text

        if not source_text.strip():
            st.warning("Please upload a file or paste some text before running the pipeline.")
        else:
            strict = masking_purpose.startswith("External Vendor")

            with st.spinner("Running PII classification & masking model..."):
                clean_text, findings = pii_model.mask(source_text, strict=strict)

            with st.container(border=True):
                st.markdown('<div class="card-heading">3. Redacted Document — Ready to Share</div>', unsafe_allow_html=True)
                st.success("✅ Document successfully masked. Send the text below, or the PDF, to any model or external recipient:")
                st.text_area("Masked Output", value=clean_text, height=200, disabled=True)

                timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                type_counts = {}
                for f in findings:
                    type_counts[f["type"]] = type_counts.get(f["type"], 0) + 1
                breakdown = (
                    "\n".join(f"- **{k}**: {v} redacted" for k, v in type_counts.items())
                    if type_counts
                    else "- No PII detected"
                )

                st.markdown("---")
                st.markdown("#### 📋 Redaction Summary")
                st.markdown(
                    f"""
- **Timestamp:** {timestamp}
- **Masking Purpose:** {masking_purpose}
- **Total Items Redacted:** {len(findings)}

**Breakdown by PII Type:**
{breakdown}
"""
                )

                pdf_bytes = build_masked_document_pdf(clean_text, timestamp, findings, masking_purpose)
                st.download_button(
                    label="📄 Download Masked Document (PDF)",
                    data=pdf_bytes,
                    file_name=f"masked_document_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf",
                    mime="application/pdf",
                )

elif st.session_state.active_page == "Prompt Engineering":
    st.write(
        "Configure a reusable prompt template on the left; watch the constructed prompt build live "
        "on the right. This page only builds the prompt — send it to Chat to actually get an answer."
    )

    config_col, results_col = st.columns(2)

    with config_col:
        with st.container(border=True):
            st.markdown('<div class="card-heading">1. Prompt Configuration</div>', unsafe_allow_html=True)
            technique = st.selectbox(
                "Select a prompt engineering strategy",
                list(PROMPT_TEMPLATES.keys()),
            )
            model_router = st.selectbox(
                "Target Backend Transformer Model",
                [
                    "Google Gemini (gemini-2.5-flash)",
                    "Groq (llama-3.3-70b-versatile)",
                ],
            )
            system_prompt = st.text_area(
                "System prompt (optional)",
                height=80,
                placeholder="e.g. You are a concise technical writer. Always answer in plain English, no jargon.",
                help="Sent separately from the user prompt via the model's system_instruction / system role.",
            )

            template = PROMPT_TEMPLATES[technique]
            st.markdown("**Template variables**")
            variable_values = {}
            for var_name, default_value in template["defaults"].items():
                variable_values[var_name] = st.text_input(
                    var_name.replace("_", " ").title(),
                    value=default_value,
                    key=f"var_{technique}_{var_name}",
                    help=f"Fills the {{{{{var_name}}}}} placeholder in the template.",
                )

            send_to_chat = st.button("💬 Send to Chat", type="primary")

    engineered_prompt = template["template"]
    for var_name, value in variable_values.items():
        engineered_prompt = engineered_prompt.replace("{{" + var_name + "}}", value)

    with results_col:
        with st.container(border=True):
            st.markdown('<div class="card-heading">2. Constructed Prompt</div>', unsafe_allow_html=True)
            st.caption("Builds live as you edit the fields on the left — no button needed to see it.")
            if system_prompt.strip():
                st.markdown("**System prompt:**")
                st.code(system_prompt, language="text")
            st.markdown("**Engineered prompt sent to the transformer model:**")
            st.code(engineered_prompt, language="text")

    if send_to_chat:
        st.session_state.chat_prefill = engineered_prompt
        st.session_state.chat_system_prompt = system_prompt
        st.session_state.chat_model_router = model_router
        st.session_state.active_page = "Chat"
        st.rerun()

elif st.session_state.active_page == "Chat":
    st.write("Chat with Gemini or Groq directly, attach a file for context, or pick up a prompt built in Prompt Engineering.")

    if "chat_history" not in st.session_state:
        st.session_state.chat_history = []
    if "chat_token_total" not in st.session_state:
        st.session_state.chat_token_total = 0

    with st.container(border=True):
        st.markdown('<div class="card-heading">Chat Configuration</div>', unsafe_allow_html=True)
        chat_col1, chat_col2 = st.columns(2)
        with chat_col1:
            default_router = st.session_state.get("chat_model_router", "Google Gemini (gemini-2.5-flash)")
            chat_model_router = st.selectbox(
                "Target Backend Transformer Model",
                ["Google Gemini (gemini-2.5-flash)", "Groq (llama-3.3-70b-versatile)"],
                index=0 if default_router.startswith("Google") else 1,
                key="chat_model_select",
            )
        with chat_col2:
            chat_system_prompt = st.text_input(
                "System prompt (optional)",
                value=st.session_state.get("chat_system_prompt", ""),
                key="chat_system_input",
            )
        chat_file = st.file_uploader(
            "Attach a file for context (optional)",
            type=["txt", "csv", "log", "md", "json", "pdf", "docx"],
            key="chat_file_uploader",
        )
        st.caption(f"🔢 Session token usage so far: {st.session_state.chat_token_total:,} total")

    st.markdown("#### 💬 Conversation")
    st.caption("Type a message in the box at the very bottom of the page to chat normally — no configuration above is required.")
    if not st.session_state.chat_history:
        st.info("No messages yet — type below to start chatting.")

    for msg in st.session_state.chat_history:
        with st.chat_message(msg["role"]):
            st.write(msg["content"])
            if msg.get("usage"):
                u = msg["usage"]
                st.caption(
                    f"Tokens — prompt: {u.get('prompt_tokens', '?')}, "
                    f"completion: {u.get('completion_tokens', '?')}, "
                    f"total: {u.get('total_tokens', '?')} · via {msg.get('model_used', '')}"
                )

    def _send_chat_message(user_text: str):
        attached_text = extract_text_from_upload(chat_file) if chat_file is not None else ""
        full_prompt = f"Document:\n{attached_text}\n\nQuestion: {user_text}" if attached_text else user_text

        st.session_state.chat_history.append({"role": "user", "content": user_text})
        if client is None and not GROQ_API_KEY:
            st.session_state.chat_history.append({
                "role": "assistant",
                "content": "Cannot generate: neither `GOOGLE_API_KEY` nor `GROQ_API_KEY` is set in the environment.",
                "usage": None,
                "model_used": None,
            })
            return
        preferred = "groq" if "Groq" in chat_model_router else "gemini"
        try:
            text, model_used, usage = generate_with_fallback(full_prompt, chat_system_prompt, preferred=preferred)
            if usage and usage.get("total_tokens"):
                st.session_state.chat_token_total += usage["total_tokens"]
            st.session_state.chat_history.append({
                "role": "assistant", "content": text, "usage": usage, "model_used": model_used,
            })
        except Exception as e:
            st.session_state.chat_history.append({
                "role": "assistant", "content": f"API Error: {e}", "usage": None, "model_used": None,
            })

    if st.session_state.get("chat_prefill"):
        prefill_text = st.session_state.pop("chat_prefill")
        with st.spinner("Sending prompt from Prompt Engineering..."):
            _send_chat_message(prefill_text)
        st.rerun()

    user_message = st.chat_input("Type your message...")
    if user_message:
        with st.spinner(f"Querying {chat_model_router}..."):
            _send_chat_message(user_message)
        st.rerun()

elif st.session_state.active_page == "Request a Model":
    st.write(
        "Don't see the model you need? Tell us which provider you'd like supported. "
        "No API key is ever typed into this app — we'll just show you the exact "
        "`.env` line to add yourself, the same safe pattern Gemini and Groq already use."
    )

    with st.container(border=True):
        st.markdown('<div class="card-heading">Request a New Model Integration</div>', unsafe_allow_html=True)
        provider_name = st.text_input(
            "Model / provider name",
            placeholder="e.g. Anthropic Claude, Mistral, Cohere",
        )
        use_case = st.text_area(
            "What would you use it for? (optional)",
            height=80,
            placeholder="e.g. Faster responses for the masking-purpose classifier",
        )
        submit_request = st.button("📨 Show Me How to Add It", type="primary")

        if submit_request:
            if not provider_name.strip():
                st.warning("Please enter a model or provider name.")
            else:
                env_var_name = re.sub(r"[^A-Za-z0-9]+", "_", provider_name.strip()).strip("_").upper() + "_API_KEY"
                st.success(f"Here's how to enable **{provider_name}**:")
                st.code(f"{env_var_name}=your_key_here", language="bash")
                st.write(
                    f"Add that line to your local `.env` file, then ask a developer to wire "
                    f"`{env_var_name}` into `generate_with_fallback()` the same way "
                    f"`GOOGLE_API_KEY` and `GROQ_API_KEY` are handled today."
                )
                if use_case.strip():
                    st.caption(f"Use case noted: {use_case.strip()}")

else:  # "How Your Data Is Masked"
    st.write(
        "🔒 Your documents never leave this machine to get masked — everything below runs "
        "locally, with no external API involved. Here's exactly what's doing the work."
    )

    with st.container(border=True):
        st.markdown('<div class="card-heading">PII Classification Model</div>', unsafe_allow_html=True)

        model_config = pii_model.nlp.model.config
        param_count = sum(p.numel() for p in pii_model.nlp.model.parameters())
        entity_types = sorted({label.split("-", 1)[-1] for label in model_config.id2label.values() if label != "O"})

        info_col1, info_col2 = st.columns(2)
        with info_col1:
            st.markdown("**Architecture**")
            st.markdown(
                f"""
- **Base model:** `bert-base-cased` ({model_config.model_type})
- **Task type:** Token classification (BIO-tagged NER)
- **Hidden size:** {model_config.hidden_size}
- **Layers:** {model_config.num_hidden_layers}
- **Attention heads:** {model_config.num_attention_heads}
- **Vocabulary size:** {model_config.vocab_size:,}
- **Total parameters:** {param_count:,} (~{param_count / 1e6:.0f}M)
- **Output labels:** {len(model_config.id2label)} BIO tags
"""
            )
        with info_col2:
            st.markdown("**Training Data**")
            st.markdown(
                """
- **Dataset:** `ai4privacy/open-pii-masking-500k-ai4privacy` (Hugging Face)
- **Language subset used:** English only (~464k training examples)
- **Fine-tuning:** 3 epochs, learning rate 2e-5, Hugging Face `Trainer`
- **Reported held-out test performance:** ~99.2% weighted accuracy / precision / recall / F1
  (token-level; dominated by the non-entity "O" class — per-type accuracy varies, see the
  notebook's classification report for the full breakdown)
"""
            )

        st.markdown("**Entity types detected** (from the 36 BIO output labels)")
        st.write(", ".join(entity_types))
